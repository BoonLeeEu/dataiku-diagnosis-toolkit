"""write_log_workbook: the sheets are laid out deterministically, every Overview link lands on its finding's row,
and input that would leak or under-explain a finding is refused before any file is touched."""

from datetime import date

import openpyxl
import pytest

import write_log_workbook as wlw


def finding(fid, level="ERROR", confidence="Likely", **kw):
    f = {
        "id": fid, "level": level, "pattern": f"[dku.x] failure {fid}", "count": 5, "first": "2026/08/01-10:00:00",
        "last": "2026/08/02-11:00:00", "first_seen_at": "run/backend.log:42",
        "excerpt": ["[2026/08/01-10:00:00.1] [t] [ERROR] [dku.x]  - failure <id>"],
        "root_cause": "The heap is too small for the workload.", "confidence": confidence,
        "check_next": ["Check install.ini for the backend Xmx"],
        "mitigation": ["1. Raise the backend Xmx", "Restart DSS"],
    }
    if confidence == "Unknown":
        f.update(root_cause="", mitigation=[])
    f.update(kw)
    return f


def log(lid="backend", findings=None, **kw):
    out = {"id": lid, "label": f"{lid}.log", "present": True,
           "files": [{"file": f"run/{lid}.log", "first": "2026/08/01-00:00:00", "last": "2026/08/03-00:00:00"}],
           "totals": {"WARN": 7, "ERROR": 5, "FATAL": 0}, "findings": findings if findings is not None else [finding("backend-001")]}
    out.update(kw)
    return out


def run(path, logs, **kw):
    args = dict(bundle="synthetic_bundle", reviewer="A. Reviewer", node_version="Design node / 14.4.3",
                diagnosis_generated="2026-08-13", logs=logs, today=date(2026, 8, 20))
    args.update(kw)
    return wlw.write_log_workbook(path, **args)


@pytest.fixture
def out(tmp_path):
    return tmp_path / "log_review.xlsx"


def test_one_sheet_per_present_log_after_the_overview(out):
    res = run(out, [log("backend"), log("nginx", [finding("nginx-001", "FATAL")]), log("hs_err", present=False, findings=[])])
    assert openpyxl.load_workbook(out).sheetnames == ["Overview", "backend.log", "nginx.log"]
    assert res["logs_not_in_bundle"] == ["hs_err"]
    assert res["by_level"] == {"FATAL": 1, "ERROR": 1, "WARN": 0}


def test_every_overview_link_lands_on_its_finding_row(out):
    run(out, [log("backend", [finding("backend-001"), finding("backend-002", "WARN", count=9)]),
              log("ipython", [finding("ipython-001", "FATAL")])])
    wb = openpyxl.load_workbook(out)
    ov, seen = wb["Overview"], 0
    for row in ov.iter_rows():
        for c in row:
            loc = c.hyperlink.location if c.hyperlink else ""
            if c.value == "Go to details":
                sheet, cell = loc.strip("'").split("'!")
                assert wb[sheet][cell].value == ov.cell(row=c.row, column=2).value
                seen += 1
    assert seen == 3


def test_overview_orders_findings_by_severity_then_count(out):
    run(out, [log("backend", [finding("backend-001", "WARN", count=99), finding("backend-002", "ERROR", count=1)]),
              log("nginx", [finding("nginx-001", "FATAL", count=2)])])
    ov = openpyxl.load_workbook(out)["Overview"]
    ids = [r[1] for r in ov.iter_rows(values_only=True) if r[1] in {"backend-001", "backend-002", "nginx-001"}]
    assert ids == ["nginx-001", "backend-002", "backend-001"]


def test_steps_are_numbered_once_and_checks_listed(out):
    run(out, [log("backend")])
    ws = openpyxl.load_workbook(out)["backend.log"]
    header = [c.value for c in ws[5]]
    row = [c.value for c in ws[6]]
    assert row[header.index("Mitigation steps")] == "1. Raise the backend Xmx\n2. Restart DSS"
    assert row[header.index("Check next, and where")] == "Check install.ini for the backend Xmx"


def test_clean_log_still_gets_a_sheet_that_says_so(out):
    res = run(out, [log("eventserver", [])])
    ws = openpyxl.load_workbook(out)["eventserver.log"]
    assert "No errors or warnings found" in ws["A6"].value
    assert res["logs_clean"] == ["eventserver"]


def test_absent_log_is_listed_on_overview_without_a_sheet(out):
    run(out, [log("backend"), log("stacks", present=False, findings=[], searched=["stacks.txt"])])
    ov = openpyxl.load_workbook(out)["Overview"]
    row = next(r for r in ov.iter_rows(values_only=True) if r[0] == "stacks.log")
    assert row[2] == "Not in bundle" and "stacks.txt" in row[7] and row[1] in (None, "")


def test_sheet_names_are_legal_unique_and_at_most_31_chars(out):
    long = "a" * 40
    run(out, [log("one", label=f"{long}/x"), log("two", label=f"{long}/x"), log("ov", label="Overview")])
    names = openpyxl.load_workbook(out).sheetnames
    assert len(set(n.lower() for n in names)) == 4 and all(len(n) <= 31 for n in names)
    assert not any(ch in n for n in names for ch in "[]:*?/\\")


@pytest.mark.parametrize("change, message", [
    ({"confidence": "Maybe"}, "confidence"),
    ({"level": "INFO"}, "level"),
    ({"count": 0}, "count"),
    ({"mitigation": []}, "needs mitigation"),
    ({"root_cause": ""}, "needs a root_cause"),
    ({"confidence": "Unknown", "root_cause": "", "mitigation": [], "check_next": []}, "check_next"),
    ({"excerpt": ["login failed for jo@example.com"]}, "an email address"),
    ({"excerpt": ["client 10.1.2.3 rejected"]}, "an IPv4 address"),
    ({"excerpt": ["password=hunter2"]}, "a credential value"),
    ({"excerpt": ["Command line: installer.sh -l License-someone-2026.json"]}, "a licence file name"),
])
def test_bad_findings_are_refused_before_any_file_is_written(out, change, message):
    with pytest.raises(wlw.LogWorkbookInputError, match=message):
        run(out, [log("backend", [finding("backend-001", **change)])])
    assert not out.exists()


def test_masked_text_is_accepted(out):
    run(out, [log("backend", [finding("backend-001", excerpt=["user=<user> from <ip> password=<secret> License-<licence>"])])])
    assert out.exists()


def test_duplicate_ids_and_wrong_extension_are_refused(out, tmp_path):
    with pytest.raises(wlw.LogWorkbookInputError, match="appears twice"):
        run(out, [log("backend", [finding("backend-001"), finding("backend-001")])])
    with pytest.raises(wlw.LogWorkbookInputError, match="appears twice"):
        run(out, [log("backend"), log("backend")])
    with pytest.raises(wlw.LogWorkbookInputError, match=".xlsx"):
        run(tmp_path / "x.csv", [log("backend")])


def test_failed_write_leaves_an_existing_file_untouched(out):
    run(out, [log("backend")])
    before = out.read_bytes()
    with pytest.raises(wlw.LogWorkbookInputError):
        run(out, [log("backend", [finding("backend-001", confidence="Maybe")])])
    assert out.read_bytes() == before
    assert not list(out.parent.glob("*.logs-tmp.xlsx"))


def test_overview_metadata_and_totals(out):
    run(out, [log("backend", totals={"WARN": 7, "ERROR": 5, "FATAL": 1})])
    ov = openpyxl.load_workbook(out)["Overview"]
    cells = {r[0]: r[1] for r in ov.iter_rows(min_row=2, max_row=6, values_only=True)}
    assert cells["Bundle:"] == "synthetic_bundle" and cells["Report generated:"] == "2026-08-20"
    row = next(r for r in ov.iter_rows(values_only=True) if r[0] == "backend.log" and r[2] in ("Findings", "Clean"))
    assert row[3:6] == (1, 5, 7) and "2026/08/01-00:00:00 to 2026/08/03-00:00:00" == row[6]
